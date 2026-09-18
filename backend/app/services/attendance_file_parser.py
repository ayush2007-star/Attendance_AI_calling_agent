from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, BinaryIO

from docx import Document
from openpyxl import Workbook, load_workbook
import pdfplumber
import xlrd

from app.services.attendance_excel_parser import (
    AttendanceImportSummary,
    parse_attendance_workbook,
)


SUPPORTED_ATTENDANCE_EXTENSIONS = {".xlsx", ".xls", ".csv", ".pdf", ".docx"}


def parse_attendance_file(
    file: BinaryIO,
    file_name: str,
) -> AttendanceImportSummary:
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_ATTENDANCE_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_ATTENDANCE_EXTENSIONS))
        raise ValueError(
            f"Unsupported attendance file format. Supported formats: {supported}"
        )

    if extension == ".xlsx":
        return parse_attendance_workbook(file)
    if extension == ".xls":
        return _parse_legacy_excel(file)
    if extension == ".csv":
        return _parse_rows(_read_csv(file))
    if extension == ".pdf":
        return _parse_rows(_read_pdf_tables(file))
    return _parse_rows(_read_docx_tables(file))


def _parse_legacy_excel(file: BinaryIO) -> AttendanceImportSummary:
    workbook = xlrd.open_workbook(file_contents=file.read())
    sheet = workbook.sheet_by_index(0)
    return _parse_rows(
        [sheet.row_values(row_number) for row_number in range(sheet.nrows)]
    )


def _read_csv(file: BinaryIO) -> list[list[Any]]:
    content = file.read().decode("utf-8-sig")
    return [list(row) for row in csv.reader(StringIO(content))]


def _read_pdf_tables(file: BinaryIO) -> list[list[Any]]:
    rows: list[list[Any]] = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    cleaned_row = [cell.strip() if isinstance(cell, str) else cell for cell in row]
                    if any(value not in (None, "") for value in cleaned_row):
                        rows.append(cleaned_row)
    return rows


def _read_docx_tables(file: BinaryIO) -> list[list[Any]]:
    document = Document(file)
    rows: list[list[Any]] = []
    for table in document.tables:
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
    return rows


def _parse_rows(rows: list[list[Any]]) -> AttendanceImportSummary:
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return parse_attendance_workbook(stream)