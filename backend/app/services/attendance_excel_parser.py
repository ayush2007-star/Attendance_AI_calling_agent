from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
import re
from typing import Any, BinaryIO

from openpyxl import load_workbook


ID_HEADER_ALIASES = {
    "student id",
    "student code",
    "roll no",
    "roll number",
    "rollno",
    "roll no / student id",
    "enrollment no",
    "enrollment number",
}
PERCENTAGE_HEADER_ALIASES = {
    "tot %",
    "tot.(%)",
    "tot (%)",
    "total %",
    "total percentage",
}
KNOWN_NON_DATE_HEADERS = {
    "s no",
    "s no.",
    "serial no",
    "name",
    "student name",
    "p",
    "a",
    "tot",
    "nf",
}
DATE_FORMATS = (
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%Y-%m-%d",
    "%d.%m.%Y",
    "%m/%d/%Y",
)


@dataclass(frozen=True)
class ParsedAttendanceRow:
    row_number: int
    student_code: str
    student_name: str | None
    attendance_by_date: dict[date, str]
    college_percentage: float | None


@dataclass
class AttendanceImportSummary:
    total_rows: int = 0
    valid_rows: int = 0
    invalid_rows: int = 0
    duplicate_rows: int = 0
    rows: list[ParsedAttendanceRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_attendance_workbook(file: BinaryIO) -> AttendanceImportSummary:
    workbook = load_workbook(file, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header_row_number, headers = _find_header_row(worksheet)
        if header_row_number is None:
            return AttendanceImportSummary(
                errors=[
                    "Could not find a header row with a student ID and date column"
                ]
            )

        id_column = _find_column(headers, ID_HEADER_ALIASES)
        percentage_column = _find_column(
            headers,
            PERCENTAGE_HEADER_ALIASES,
        )
        date_columns, invalid_date_headers = _find_date_columns(headers)

        summary = AttendanceImportSummary()
        summary.warnings.extend(
            f"Unrecognized date-like header: {header}"
            for header in invalid_date_headers
        )

        if id_column is None:
            summary.errors.append("Missing student ID or roll number column")
        if not date_columns:
            summary.errors.append("No valid attendance date columns found")
        if summary.errors:
            return summary

        seen_student_codes: set[str] = set()
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=header_row_number + 1,
                values_only=True,
            ),
            start=header_row_number + 1,
        ):
            if _is_blank_row(values):
                continue

            summary.total_rows += 1
            row_errors: list[str] = []
            raw_student_code = values[id_column]
            student_code = _clean_text(raw_student_code)
            if not student_code:
                row_errors.append("missing student ID")

            if student_code and student_code in seen_student_codes:
                summary.duplicate_rows += 1
                row_errors.append(f"duplicate student ID: {student_code}")
            elif student_code:
                seen_student_codes.add(student_code)

            percentage = None
            if percentage_column is not None:
                percentage, percentage_error = _parse_percentage(
                    values[percentage_column]
                )
                if percentage_error:
                    row_errors.append(percentage_error)
            else:
                summary.warnings.append(
                    "Tot.(%) column was not provided; percentage fallback is required"
                )

            attendance_by_date: dict[date, str] = {}
            for attendance_date, column_index in date_columns.items():
                value = _clean_text(values[column_index])
                if value:
                    attendance_by_date[attendance_date] = value

            if not attendance_by_date:
                row_errors.append("no date-wise attendance values")

            if row_errors:
                summary.invalid_rows += 1
                summary.errors.extend(
                    f"row {row_number}: {error}" for error in row_errors
                )
                continue

            student_name = _find_row_name(values, headers)
            summary.rows.append(
                ParsedAttendanceRow(
                    row_number=row_number,
                    student_code=student_code,
                    student_name=student_name,
                    attendance_by_date=attendance_by_date,
                    college_percentage=percentage,
                )
            )
            summary.valid_rows += 1

        return summary
    finally:
        workbook.close()


def _find_header_row(worksheet: Any) -> tuple[int | None, list[str]]:
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=1, max_row=20, values_only=True),
        start=1,
    ):
        headers = [_normalize_header(value) for value in values]
        has_id = any(header in ID_HEADER_ALIASES for header in headers)
        has_date = any(_parse_date(value) is not None for value in values)
        if has_id and has_date:
            return row_number, headers
    return None, []


def _find_column(headers: list[str], aliases: set[str]) -> int | None:
    for index, header in enumerate(headers):
        if header in aliases:
            return index
    return None


def _find_date_columns(
    headers: list[str],
) -> tuple[dict[date, int], list[str]]:
    date_columns: dict[date, int] = {}
    invalid_date_headers: list[str] = []
    for index, header in enumerate(headers):
        parsed_date = _parse_date(header)
        if parsed_date is not None:
            if parsed_date in date_columns:
                invalid_date_headers.append(header)
            else:
                date_columns[parsed_date] = index
        elif _looks_like_date_header(header) and header not in KNOWN_NON_DATE_HEADERS:
            invalid_date_headers.append(header)
    return date_columns, invalid_date_headers


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None

    text = value.strip()
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        pass

    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def _parse_percentage(value: Any) -> tuple[float | None, str | None]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing total attendance percentage"

    if isinstance(value, str):
        text = value.strip().replace("%", "")
        try:
            percentage = float(text)
        except ValueError:
            return None, f"invalid attendance percentage: {value}"
    elif isinstance(value, int | float):
        percentage = float(value)
    else:
        return None, f"invalid attendance percentage: {value}"

    if not 0 <= percentage <= 100:
        return None, f"attendance percentage out of range: {percentage}"
    return percentage, None


def _find_row_name(values: tuple[Any, ...], headers: list[str]) -> str | None:
    for index, header in enumerate(headers):
        if header in {"name", "student name"}:
            return _clean_text(values[index]) or None
    return None


def _normalize_header(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    text = _clean_text(value).lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _looks_like_date_header(value: str) -> bool:
    return bool(re.search(r"\d", value)) and bool(
        re.search(r"[-/.]", value)
    )


def _is_blank_row(values: tuple[Any, ...]) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)
